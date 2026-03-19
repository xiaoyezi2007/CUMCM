import openpyxl

wb = openpyxl.load_workbook('result1.xlsx')
sheet = wb['速度']
for i in [2,3,53,103,153,203,224]:
    print(sheet.cell(i, 1).value,end = ' &')
    for t in [2,62,122,182,242,302]:
        print(sheet.cell(i, t).value,end = ' &')
    print('\\\\ \\hline')
