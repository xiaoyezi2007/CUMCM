from sympy import *
import openpyxl

l_head = 286
l_body = 165
a = 27.5 / pi
r0 = 16 * 55
v0 = 100


def ans_site(t):
    ansr = []
    anssita = []
    anssite = []
    r = symbols('r')
    fr = r * sqrt(r * r + a * a) / 2 + a * a * ln(
        (r + sqrt(r * r + a * a)) / (r0 + sqrt(r0 * r0 + a * a))) / 2 - r0 * sqrt(
        r0 * r0 + a * a) / 2 + a * v0 * t
    ansr.append(nsolve(fr, 700))
    anssita.append(ansr[0] / a)
    x0 = ansr[0] * cos(anssita[0])
    y0 = ansr[0] * sin(anssita[0])
    anssite.append(x0.evalf())
    anssite.append(y0.evalf())
    fr = (r * cos(r / a) - x0) * (r * cos(r / a) - x0) + (r * sin(r / a) - y0) * (r * sin(r / a) - y0) - l_head * l_head
    ansr.append(nsolve(fr, ansr[0] + 2))
    anssita.append(ansr[1] / a)
    for i in range(1, 224):
        x0 = ansr[i] * cos(anssita[i])
        y0 = ansr[i] * sin(anssita[i])
        fr = (r * cos(r / a) - x0) * (r * cos(r / a) - x0) + (r * sin(r / a) - y0) * (
                r * sin(r / a) - y0) - l_body * l_body
        ansr.append(nsolve(fr, ansr[i] + 2))
        anssita.append(ansr[i + 1] / a)
        x = ansr[i] * cos(anssita[i])
        y = ansr[i] * sin(anssita[i])
        anssite.append(x.evalf())
        anssite.append(y.evalf())
    return anssite


wb = openpyxl.load_workbook('result2.xlsx')
sheet = wb['Sheet1']
ans = ans_site(412.473837)
for i in range(224):
    tmp1 = ans[2 * i] / 100
    tmp2 = ans[2 * i + 1] / 100
    sheet.cell(i + 2, 2).value = str(round(tmp1, 6))
    sheet.cell(i + 2, 3).value = str(round(tmp2, 6))
wb.save('result2.xlsx')
